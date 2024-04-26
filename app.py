import os
import re
from flask import Flask, render_template, request, send_file
from werkzeug.utils import secure_filename
from PyPDF2 import PdfReader
from docx import Document
from openpyxl import Workbook

app = Flask(__name__)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {'pdf', 'doc', 'docx'}

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

def allowed_file(filename):
    # print(filename.rsplit('.', 1))
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def extract_information_pdf(pdf_path):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'(\d{3}[-\.\s]??\d{3}[-\.\s]??\d{4}|\(\d{3}\)\s*\d{3}[-\.\s]??\d{4}|\d{3}[-\.\s]??\d{4})'
    
    text = ""
    emails = []
    phones = []
    
    with open(pdf_path, 'rb') as f:
        reader = PdfReader(f)
        for page in reader.pages:
            text += page.extract_text()
    
    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    
    return emails, phones, text

def extract_information_docx(docx_path):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    phone_pattern = r'(\d{3}[-\.\s]??\d{3}[-\.\s]??\d{4}|\(\d{3}\)\s*\d{3}[-\.\s]??\d{4}|\d{3}[-\.\s]??\d{4})'
    
    text = ""
    emails = []
    phones = []
    
    doc = Document(docx_path)
    for para in doc.paragraphs:
        text += para.text
    
    emails = re.findall(email_pattern, text)
    phones = re.findall(phone_pattern, text)
    
    return emails, phones, text

@app.route('/')
def index():
    return render_template('upload.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return 'No file part'
    
    files = request.files.getlist('file')
    
    if len(files) == 0:
        return 'No files selected'
    
    wb = Workbook()
    ws = wb.active
    ws.append(['Email', 'Phone', 'Text'])
    
    for file in files:
        if file and allowed_file(file.filename):
            filename = secure_filename(file.filename)
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
            file.save(file_path)
            
            if filename.endswith('.pdf'):
                emails, phones, text = extract_information_pdf(file_path)
            elif filename.endswith('.doc') or filename.endswith('.docx'):
                emails, phones, text = extract_information_docx(file_path)
            else:
                return 'Unsupported file format'
            
            for email, phone in zip(emails, phones):
                ws.append([email, phone, text])
            
    excel_filename = 'all_data.xlsx'
    excel_file_path = os.path.join(app.config['UPLOAD_FOLDER'], excel_filename)
    wb.save(excel_file_path)
    
    return send_file(excel_file_path, as_attachment=True)

if __name__ == '__main__':
    app.run(debug=True)
